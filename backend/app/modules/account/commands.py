"""账户管理写库与查询。广告主同步走 oceanengine.sync，授权链接走 oceanengine.oauth。"""

from __future__ import annotations

import hashlib
import hmac
import io
import json
from typing import Any
from urllib.parse import unquote

from openpyxl import load_workbook

from sqlalchemy import delete, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.envelope import ApiError
from app.core.pagination import PageParams, page_data
from app.core.times import beijing_now
from app.modules.account.model import (
    AccountSyncRun,
    AdvertiserAccount,
    AdvertiserNameLog,
    DeliverySubject,
    DouyinAccount,
    DouyinPitcher,
    OeApp,
    OeOrganization,
    OePromotion,
    ProductLibrary,
    ProductLibraryPitcher,
)
from app.modules.account.schemas import (
    AssignPitchersBody,
    DouyinBody,
    DouyinQuery,
    DouyinUpdateBody,
    ProductLibraryBody,
    ProductLibraryQuery,
    RenameAdvertiserItem,
    SubjectBody,
    SubjectQuery,
)
from app.modules.system_admin.domain.models import Department, User, UserTag, UserTagLink

_ASSIGNED_PITCHER = "广告主已有投手，须先解绑"
_PITCHER_TAG = "投手"
_MISSING_FALLBACK = "缺少兜底库"



def _oauth_state_body(payload: dict[str, Any]) -> str:
    body = {key: value for key, value in payload.items() if key != "sig"}
    return json.dumps(body, ensure_ascii=False, separators=(",", ":"), sort_keys=True)


def sign_oauth_state(payload: dict[str, Any]) -> str:
    """授权 state：紧凑 JSON，多一个 sig。HMAC-SHA256 的密钥是 jwt_secret。

    正文必须带整数 appRowId（oe_app.id），并保留 agentId；自研再带 agency=true。
    这样 oceanengine.oauth 仍能按 JSON 认出渠道。拼进 URL 时把双引号写成 %22。
    """
    raw = _oauth_state_body(payload)
    signature = hmac.new(get_settings().jwt_secret.encode(), raw.encode(), hashlib.sha256).hexdigest()
    signed = json.loads(raw)
    signed["sig"] = signature
    return json.dumps(signed, ensure_ascii=False, separators=(",", ":"), sort_keys=True)


def _decode_oauth_state(state: str) -> dict[str, Any]:
    try:
        payload = json.loads(unquote(state).strip())
    except json.JSONDecodeError:
        raise ApiError(422, "state 校验失败") from None
    if not isinstance(payload, dict):
        raise ApiError(422, "state 校验失败")
    signature = payload.get("sig")
    if not isinstance(signature, str) or not signature:
        raise ApiError(422, "state 校验失败")
    expected = hmac.new(
        get_settings().jwt_secret.encode(),
        _oauth_state_body(payload).encode(),
        hashlib.sha256,
    ).hexdigest()
    if not hmac.compare_digest(expected, signature):
        raise ApiError(422, "state 校验失败")
    return payload


async def assert_oauth_state(session: AsyncSession, state: str) -> None:
    """非空 state 必须通过 HMAC，且 appRowId 对得上未删除的应用行。"""
    payload = _decode_oauth_state(state)
    app_row_id = payload.get("appRowId")
    if isinstance(app_row_id, bool) or not isinstance(app_row_id, int):
        raise ApiError(422, "state 校验失败")
    app = await session.get(OeApp, app_row_id)
    if app is None or app.is_deleted:
        raise ApiError(422, "state 校验失败")
    agent_id = payload.get("agentId")
    if agent_id is not None and str(agent_id) != app.agent_key:
        raise ApiError(422, "state 校验失败")


def _ilike_pattern(keyword: str) -> str:
    escaped = keyword.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
    return f"%{escaped}%"


def _promotion_label(row: OePromotion) -> str:
    name = (row.name or "").strip()
    return name or str(row.promotion_id)


def _busy_message(rows: list[OePromotion]) -> str:
    ordered = sorted(rows, key=lambda item: item.promotion_id)
    return "存在执行中的广告：" + "、".join(_promotion_label(row) for row in ordered)


def _money(value: Any) -> float:
    return float(value)


def _money_or_none(value: Any) -> float | None:
    if value is None:
        return None
    return float(value)


async def _require_pitchers(session: AsyncSession, user_ids: list[int]) -> None:
    unique = list(dict.fromkeys(user_ids))
    if not unique:
        return
    result = await session.execute(
        select(User.id)
        .join(UserTagLink, UserTagLink.user_id == User.id)
        .join(UserTag, UserTag.id == UserTagLink.tag_id)
        .where(
            User.id.in_(unique),
            User.is_deleted == 0,
            User.enabled.is_(True),
            UserTag.is_deleted == 0,
            UserTag.name == _PITCHER_TAG,
            UserTagLink.is_deleted == 0,
        )
    )
    found = set(result.scalars().all())
    if found != set(unique):
        raise ApiError(422, "投手不可用")


async def _alive_advertisers(
    session: AsyncSession, ocean_ids: list[int]
) -> dict[int, AdvertiserAccount]:
    unique = list(dict.fromkeys(ocean_ids))
    result = await session.execute(
        select(AdvertiserAccount).where(
            AdvertiserAccount.advertiser_id.in_(unique),
            AdvertiserAccount.is_deleted == 0,
        )
    )
    found = {row.advertiser_id: row for row in result.scalars().all()}
    if any(item not in found for item in unique):
        raise ApiError(422, "广告主不存在或已解绑")
    return found


def _write_assign(row: AdvertiserAccount, pitcher_user_id: int | None, operator: str) -> None:
    """写入投手。pitcher_user_id 为空表示取消分配。"""
    row.pitcher_user_id = pitcher_user_id
    row.assigned_at = beijing_now() if pitcher_user_id is not None else None
    row.assigned_by = operator if pitcher_user_id is not None else None


async def _enabled_promotions(
    session: AsyncSession, *, advertiser_ids: list[int] | None = None, douyin_account_id: int | None = None
) -> list[OePromotion]:
    filters = [OePromotion.opt_status == "ENABLE", OePromotion.is_deleted == 0]
    if advertiser_ids is not None:
        filters.append(OePromotion.advertiser_id.in_(advertiser_ids))
    if douyin_account_id is not None:
        filters.append(OePromotion.douyin_account_id == douyin_account_id)
    result = await session.execute(select(OePromotion).where(*filters))
    return list(result.scalars().all())


async def _reject_if_busy(rows: list[OePromotion]) -> None:
    if rows:
        raise ApiError(409, _busy_message(rows))


async def assign_advertisers(
    session: AsyncSession,
    advertiser_ids: list[int],
    pitcher_user_id: int,
    operator: str,
) -> dict[str, Any]:
    """任一户已有投手则整批 409，不部分写入。"""
    accounts = await _alive_advertisers(session, advertiser_ids)
    if any(row.pitcher_user_id is not None for row in accounts.values()):
        raise ApiError(409, _ASSIGNED_PITCHER)
    await _require_pitchers(session, [pitcher_user_id])
    for row in accounts.values():
        _write_assign(row, pitcher_user_id, operator)
    await session.commit()
    return {"advertiser_ids": advertiser_ids, "pitcher_user_id": pitcher_user_id}


async def rename_advertisers(
    session: AsyncSession,
    items: list[RenameAdvertiserItem],
    operator: str,
) -> dict[str, Any]:
    """写 local_name，并给每个户记一行改名审计。"""
    ocean_ids = [item.advertiser_id for item in items]
    if len(ocean_ids) != len(set(ocean_ids)):
        raise ApiError(422, "广告主 id 重复")
    accounts = await _alive_advertisers(session, ocean_ids)
    rows: list[dict[str, Any]] = []
    for item in items:
        account = accounts[item.advertiser_id]
        old_name = account.local_name or account.name
        account.local_name = item.name
        session.add(
            AdvertiserNameLog(
                advertiser_account_id=account.id,
                old_name=old_name,
                new_name=item.name,
                operator_account=operator,
            )
        )
        rows.append({"advertiser_id": item.advertiser_id, "account_name": item.name})
    await session.commit()
    return {"list": rows}


async def unbind_advertisers(session: AsyncSession, advertiser_ids: list[int]) -> None:
    """有执行中广告则整批 409。解绑是软删，投手和本地名留在已删行上。"""
    accounts = await _alive_advertisers(session, advertiser_ids)
    ocean_ids = list(accounts)
    busy = await _enabled_promotions(session, advertiser_ids=ocean_ids)
    await _reject_if_busy(busy)
    now = beijing_now()
    for row in accounts.values():
        row.mark_deleted()
        row.unbound_at = now
    await session.commit()


def _cell_text(value: Any) -> str:
    """单元格转成去掉首尾空白的文本。整数和整数值的浮点按十进制写出，避免科学计数。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        raise ApiError(422, "广告主 id 不是整数")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not value.is_integer():
            raise ApiError(422, "广告主 id 不是整数")
        return str(int(value))
    return str(value).strip()


def _xlsx_rows(raw: bytes, action: str) -> list[dict[str, str]]:
    """读第一张表。表头必须与 action 要求的列完全一致，空行跳过。"""
    if not raw.startswith(b"PK"):
        raise ApiError(422, "文件不是 xlsx")
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ApiError(422, "文件不是 xlsx") from None
    try:
        sheet = workbook.active
        if sheet is None:
            raise ApiError(422, "文件不是 xlsx")
        row_iter = sheet.iter_rows(values_only=True)
        header_row = next(row_iter, None)
        if header_row is None:
            raise ApiError(422, "表头不正确")
        headers = [_cell_text(cell) for cell in header_row]
        while headers and headers[-1] == "":
            headers.pop()
        expected = ["广告主 id", "投手登录账号"] if action == "assign" else ["广告主 id"]
        if headers != expected:
            raise ApiError(422, "表头不正确")
        parsed: list[dict[str, str]] = []
        for record in row_iter:
            values = {
                headers[index]: _cell_text(record[index] if index < len(record) else None)
                for index in range(len(headers))
            }
            if not any(values.values()):
                continue
            parsed.append(values)
        if not parsed:
            raise ApiError(422, "文件没有数据行")
        return parsed
    finally:
        workbook.close()


def _parse_ocean_id(text: str) -> int:
    try:
        return int(text)
    except ValueError:
        raise ApiError(422, "广告主 id 不是整数") from None


async def _user_by_login_pitcher(session: AsyncSession, login_account: str) -> User:
    result = await session.execute(
        select(User).where(User.login_account == login_account, User.is_deleted == 0)
    )
    user = result.scalar_one_or_none()
    if user is None or not user.enabled:
        raise ApiError(422, "投手不可用")
    await _require_pitchers(session, [user.id])
    return user


async def import_advertisers(
    session: AsyncSession,
    action: str,
    raw: bytes,
    operator: str,
) -> dict[str, Any]:
    """xlsx 批量分配或解绑，拒绝规则与单批接口相同，整份失败不写。"""
    if action not in ("assign", "unbind"):
        raise ApiError(422, "action 只允许 assign 或 unbind")
    records = _xlsx_rows(raw, action)
    if action == "unbind":
        ocean_ids = [_parse_ocean_id(row["广告主 id"]) for row in records]
        await unbind_advertisers(session, ocean_ids)
        return {"action": "unbind", "success_count": len(records)}

    pairs: list[tuple[int, int]] = []
    seen: dict[int, str] = {}
    for row in records:
        ocean_id = _parse_ocean_id(row["广告主 id"])
        login = row["投手登录账号"]
        if not login:
            raise ApiError(422, "投手不可用")
        previous = seen.get(ocean_id)
        if previous is not None and previous != login:
            raise ApiError(422, "同一广告主不能分给多个投手")
        seen[ocean_id] = login
        user = await _user_by_login_pitcher(session, login)
        pairs.append((ocean_id, user.id))
    accounts = await _alive_advertisers(session, [ocean_id for ocean_id, _user_id in pairs])
    if any(row.pitcher_user_id is not None for row in accounts.values()):
        raise ApiError(409, _ASSIGNED_PITCHER)
    for ocean_id, pitcher_user_id in pairs:
        _write_assign(accounts[ocean_id], pitcher_user_id, operator)
    await session.commit()
    return {"action": "assign", "success_count": len(records)}


def _subject_data(row: DeliverySubject, ocean_advertiser_id: int) -> dict[str, Any]:
    return {
        "id": row.id,
        "name": row.name,
        "subject_no": row.subject_no,
        "short_name": row.short_name,
        "delivery_mode": row.delivery_mode,
        "theater_name": row.theater_name,
        "theater_kind": row.theater_kind,
        "charge_mode": row.charge_mode,
        "min_bid": _money(row.min_bid),
        "max_bid": _money(row.max_bid),
        "roi_goal": _money_or_none(row.roi_goal),
        "department_id": row.department_id,
        "material_account_id": ocean_advertiser_id,
        "dual_bid": row.dual_bid,
        "bid_panel": row.bid_panel,
    }


async def _require_department(session: AsyncSession, department_id: int | None) -> None:
    if department_id is None:
        return
    department = await session.get(Department, department_id)
    if department is None or department.is_deleted:
        raise ApiError(422, "部门不存在")


async def _require_owner(session: AsyncSession, user_id: int | None) -> None:
    if user_id is None:
        return
    user = await session.get(User, user_id)
    if user is None or user.is_deleted:
        raise ApiError(422, "负责人不存在")


async def _require_material_account(session: AsyncSession, ocean_advertiser_id: int) -> AdvertiserAccount:
    """素材账户按广告主列表里的 advertiser_id 查，不按本库自增主键。"""
    account = await session.scalar(
        select(AdvertiserAccount).where(
            AdvertiserAccount.advertiser_id == ocean_advertiser_id,
            AdvertiserAccount.is_deleted == 0,
        )
    )
    if account is None:
        raise ApiError(422, "素材账户不存在或已解绑")
    return account


async def _subject_no_free(session: AsyncSession, subject_no: int, exclude_id: int | None) -> None:
    stmt = select(DeliverySubject.id).where(
        DeliverySubject.subject_no == subject_no,
        DeliverySubject.is_deleted == 0,
    )
    if exclude_id is not None:
        stmt = stmt.where(DeliverySubject.id != exclude_id)
    if await session.scalar(stmt) is not None:
        raise ApiError(409, "主体 id 已存在")


def _apply_subject(row: DeliverySubject, body: SubjectBody, material_account_pk: int) -> None:
    row.name = body.name
    row.subject_no = body.subject_no
    row.short_name = body.short_name
    row.delivery_mode = body.delivery_mode
    row.theater_name = body.theater_name
    row.theater_kind = body.theater_kind
    row.charge_mode = body.charge_mode
    row.min_bid = body.min_bid
    row.max_bid = body.max_bid
    row.roi_goal = body.roi_goal
    row.department_id = body.department_id
    row.material_account_id = material_account_pk
    row.dual_bid = body.dual_bid
    row.bid_panel = body.bid_panel


async def list_subjects(session: AsyncSession, query: SubjectQuery) -> dict[str, Any]:
    filters = [DeliverySubject.is_deleted == 0]
    if query.name:
        filters.append(DeliverySubject.name.ilike(_ilike_pattern(query.name), escape="\\"))
    if query.subject_no is not None:
        filters.append(DeliverySubject.subject_no == query.subject_no)
    if query.delivery_mode is not None:
        filters.append(DeliverySubject.delivery_mode == query.delivery_mode)
    if query.theater_kind:
        filters.append(DeliverySubject.theater_kind == query.theater_kind)
    total = int(await session.scalar(select(func.count()).select_from(DeliverySubject).where(*filters)) or 0)
    params = PageParams(page=query.page, page_size=query.page_size)
    result = await session.execute(
        select(DeliverySubject)
        .where(*filters)
        .order_by(DeliverySubject.id.desc())
        .offset(params.offset)
        .limit(params.page_size)
    )
    rows = list(result.scalars().all())
    ocean_ids = await _ocean_advertiser_ids(session, [row.material_account_id for row in rows])
    items = [_subject_data(row, ocean_ids.get(row.material_account_id, 0)) for row in rows]
    return page_data(items, total, params)


async def _ocean_advertiser_ids(session: AsyncSession, account_pks: list[int]) -> dict[int, int]:
    if not account_pks:
        return {}
    result = await session.execute(
        select(AdvertiserAccount.id, AdvertiserAccount.advertiser_id).where(
            AdvertiserAccount.id.in_(account_pks)
        )
    )
    return {pk: ocean_id for pk, ocean_id in result.all()}


async def create_subject(session: AsyncSession, body: SubjectBody) -> dict[str, Any]:
    await _subject_no_free(session, body.subject_no, None)
    await _require_department(session, body.department_id)
    account = await _require_material_account(session, body.material_account_id)
    row = DeliverySubject(
        name=body.name,
        subject_no=body.subject_no,
        short_name=body.short_name,
        delivery_mode=body.delivery_mode,
        theater_name=body.theater_name,
        theater_kind=body.theater_kind,
        charge_mode=body.charge_mode,
        min_bid=body.min_bid,
        max_bid=body.max_bid,
        roi_goal=body.roi_goal,
        department_id=body.department_id,
        material_account_id=account.id,
        dual_bid=body.dual_bid,
        bid_panel=body.bid_panel,
    )
    session.add(row)
    await session.flush()
    data = _subject_data(row, account.advertiser_id)
    await session.commit()
    return data


async def update_subject(session: AsyncSession, subject_id: int, body: SubjectBody) -> dict[str, Any]:
    row = await session.get(DeliverySubject, subject_id)
    if row is None or row.is_deleted:
        raise ApiError(404, "投放主体不存在")
    await _subject_no_free(session, body.subject_no, row.id)
    await _require_department(session, body.department_id)
    account = await _require_material_account(session, body.material_account_id)
    _apply_subject(row, body, account.id)
    data = _subject_data(row, account.advertiser_id)
    await session.commit()
    return data


async def _pitcher_ids_by_douyin(session: AsyncSession, account_ids: list[int]) -> dict[int, list[int]]:
    if not account_ids:
        return {}
    result = await session.execute(
        select(DouyinPitcher.douyin_account_id, DouyinPitcher.user_id)
        .where(DouyinPitcher.douyin_account_id.in_(account_ids))
        .order_by(DouyinPitcher.user_id)
    )
    grouped: dict[int, list[int]] = {item: [] for item in account_ids}
    for account_id, user_id in result.all():
        grouped.setdefault(account_id, []).append(user_id)
    return grouped


def _douyin_data(row: DouyinAccount, pitcher_user_ids: list[int]) -> dict[str, Any]:
    pitchers = [] if row.delivery_mode == "standard" else pitcher_user_ids
    return {
        "id": row.id,
        "aweme_id": row.aweme_id,
        "name": row.name,
        "delivery_mode": row.delivery_mode,
        "enabled": row.enabled,
        "department_id": None if row.delivery_mode == "standard" else row.department_id,
        "owner_user_id": None if row.delivery_mode == "standard" else row.owner_user_id,
        "pitcher_user_ids": pitchers,
    }


async def _aweme_free(
    session: AsyncSession, delivery_mode: str, aweme_id: str, exclude_id: int | None
) -> None:
    stmt = select(DouyinAccount.id).where(
        DouyinAccount.delivery_mode == delivery_mode,
        DouyinAccount.aweme_id == aweme_id,
        DouyinAccount.is_deleted == 0,
    )
    if exclude_id is not None:
        stmt = stmt.where(DouyinAccount.id != exclude_id)
    if await session.scalar(stmt) is not None:
        raise ApiError(409, "抖音号已存在")


async def _douyin_pitcher_count(session: AsyncSession, account_id: int) -> int:
    total = await session.scalar(
        select(func.count())
        .select_from(DouyinPitcher)
        .where(DouyinPitcher.douyin_account_id == account_id)
    )
    return int(total or 0)


async def _alive_douyin(session: AsyncSession, account_id: int) -> DouyinAccount:
    row = await session.get(DouyinAccount, account_id)
    if row is None or row.is_deleted:
        raise ApiError(404, "抖音号不存在")
    return row


async def list_douyin(session: AsyncSession, query: DouyinQuery) -> dict[str, Any]:
    filters = [
        DouyinAccount.is_deleted == 0,
        DouyinAccount.delivery_mode == query.delivery_mode,
    ]
    if query.aweme_id:
        filters.append(DouyinAccount.aweme_id == query.aweme_id)
    if query.name:
        filters.append(DouyinAccount.name.ilike(_ilike_pattern(query.name), escape="\\"))
    if query.department_id is not None:
        filters.append(DouyinAccount.department_id == query.department_id)
    total = int(await session.scalar(select(func.count()).select_from(DouyinAccount).where(*filters)) or 0)
    params = PageParams(page=query.page, page_size=query.page_size)
    result = await session.execute(
        select(DouyinAccount)
        .where(*filters)
        .order_by(DouyinAccount.id.desc())
        .offset(params.offset)
        .limit(params.page_size)
    )
    rows = list(result.scalars().all())
    pitchers = await _pitcher_ids_by_douyin(session, [row.id for row in rows])
    items = [_douyin_data(row, pitchers.get(row.id, [])) for row in rows]
    return page_data(items, total, params)


async def create_douyin(session: AsyncSession, body: DouyinBody, operator: str) -> dict[str, Any]:
    await _aweme_free(session, body.delivery_mode, body.aweme_id, None)
    await _require_department(session, body.department_id)
    await _require_owner(session, body.owner_user_id)
    row = DouyinAccount(
        aweme_id=body.aweme_id,
        name=body.name,
        delivery_mode=body.delivery_mode,
        enabled=body.enabled,
        department_id=body.department_id,
        owner_user_id=body.owner_user_id,
        created_by=operator,
    )
    session.add(row)
    await session.flush()
    data = _douyin_data(row, [])
    await session.commit()
    return data


async def update_douyin(session: AsyncSession, account_id: int, body: DouyinUpdateBody) -> dict[str, Any]:
    row = await _alive_douyin(session, account_id)
    await _aweme_free(session, body.delivery_mode, body.aweme_id, row.id)
    await _require_department(session, body.department_id)
    await _require_owner(session, body.owner_user_id)
    if body.delivery_mode == "standard" and await _douyin_pitcher_count(session, row.id):
        raise ApiError(422, "标准号不能分配投手，须先处理分配后再改模式")
    if row.enabled and not body.enabled:
        await _reject_if_busy(await _enabled_promotions(session, douyin_account_id=row.id))
    row.aweme_id = body.aweme_id
    row.name = body.name
    row.delivery_mode = body.delivery_mode
    row.enabled = body.enabled
    row.department_id = body.department_id
    row.owner_user_id = body.owner_user_id
    pitchers = await _pitcher_ids_by_douyin(session, [row.id])
    data = _douyin_data(row, pitchers.get(row.id, []))
    await session.commit()
    return data


async def set_douyin_enabled(
    session: AsyncSession, account_id: int, enabled: bool
) -> dict[str, Any]:
    row = await _alive_douyin(session, account_id)
    if row.enabled and not enabled:
        await _reject_if_busy(await _enabled_promotions(session, douyin_account_id=row.id))
    row.enabled = enabled
    await session.commit()
    return {"id": row.id, "enabled": row.enabled}


async def assign_douyin(
    session: AsyncSession, account_id: int, body: AssignPitchersBody
) -> dict[str, Any]:
    """覆盖这一个全域号的投手。标准号 422。不改部门和负责人。"""
    row = await _alive_douyin(session, account_id)
    if row.delivery_mode != "uni":
        raise ApiError(422, "标准号不能分配投手")
    user_ids = list(dict.fromkeys(body.pitcher_user_ids))
    await _require_pitchers(session, user_ids)
    await session.execute(delete(DouyinPitcher).where(DouyinPitcher.douyin_account_id == row.id))
    await session.flush()
    for user_id in user_ids:
        session.add(DouyinPitcher(douyin_account_id=row.id, user_id=user_id))
    await session.commit()
    return {"id": row.id, "pitcher_user_ids": user_ids}


async def reclaim_douyin(session: AsyncSession, account_id: int) -> dict[str, Any]:
    """只清部门和负责人。有执行中广告则整次 409，投手分配保留。"""
    row = await _alive_douyin(session, account_id)
    if row.delivery_mode != "uni":
        raise ApiError(422, "标准号没有部门与负责人可回收")
    await _reject_if_busy(await _enabled_promotions(session, douyin_account_id=row.id))
    row.department_id = None
    row.owner_user_id = None
    pitchers = await _pitcher_ids_by_douyin(session, [row.id])
    await session.commit()
    return {
        "id": row.id,
        "department_id": None,
        "owner_user_id": None,
        "pitcher_user_ids": pitchers.get(row.id, []),
    }


async def delete_douyin(session: AsyncSession, account_id: int) -> None:
    row = await _alive_douyin(session, account_id)
    await _reject_if_busy(await _enabled_promotions(session, douyin_account_id=row.id))
    row.mark_deleted()
    await session.commit()


async def _org_by_ocean_id(session: AsyncSession, ocean_account_id: int) -> OeOrganization:
    result = await session.execute(
        select(OeOrganization).where(
            OeOrganization.ocean_account_id == ocean_account_id,
            OeOrganization.is_deleted == 0,
        )
    )
    rows = list(result.scalars().all())
    if len(rows) != 1:
        raise ApiError(422, "组织不存在")
    return rows[0]


async def _library_pitchers(session: AsyncSession, library_ids: list[int]) -> dict[int, list[int]]:
    if not library_ids:
        return {}
    result = await session.execute(
        select(ProductLibraryPitcher.product_library_id, ProductLibraryPitcher.user_id)
        .where(ProductLibraryPitcher.product_library_id.in_(library_ids))
        .order_by(ProductLibraryPitcher.user_id)
    )
    grouped: dict[int, list[int]] = {item: [] for item in library_ids}
    for library_id, user_id in result.all():
        grouped.setdefault(library_id, []).append(user_id)
    return grouped


def _library_data(row: ProductLibrary, ocean_account_id: int, pitcher_user_ids: list[int]) -> dict[str, Any]:
    return {
        "id": row.id,
        "name": row.name,
        "library_no": row.library_no,
        "library_kind": row.library_kind,
        "organization_id": ocean_account_id,
        "library_role": row.library_role,
        "uploaded_count": row.uploaded_count,
        "pitcher_user_ids": [] if row.library_role == "fallback" else pitcher_user_ids,
    }


async def _library_no_free(session: AsyncSession, library_no: int, exclude_id: int | None) -> None:
    stmt = select(ProductLibrary.id).where(
        ProductLibrary.library_no == library_no,
        ProductLibrary.is_deleted == 0,
    )
    if exclude_id is not None:
        stmt = stmt.where(ProductLibrary.id != exclude_id)
    if await session.scalar(stmt) is not None:
        raise ApiError(409, "商品库 id 已存在")


async def _reject_second_fallback(
    session: AsyncSession, organization_id: int, library_kind: str, exclude_id: int | None
) -> None:
    stmt = select(ProductLibrary.id).where(
        ProductLibrary.organization_id == organization_id,
        ProductLibrary.library_kind == library_kind,
        ProductLibrary.library_role == "fallback",
        ProductLibrary.is_deleted == 0,
    )
    if exclude_id is not None:
        stmt = stmt.where(ProductLibrary.id != exclude_id)
    if await session.scalar(stmt) is not None:
        raise ApiError(409, "同一组织同一类型已有兜底库")


async def _alive_library(session: AsyncSession, library_id: int) -> ProductLibrary:
    row = await session.get(ProductLibrary, library_id)
    if row is None or row.is_deleted:
        raise ApiError(404, "商品库不存在")
    return row


async def list_product_libraries(session: AsyncSession, query: ProductLibraryQuery) -> dict[str, Any]:
    filters = [ProductLibrary.is_deleted == 0]
    if query.name:
        filters.append(ProductLibrary.name.ilike(_ilike_pattern(query.name), escape="\\"))
    if query.library_kind is not None:
        filters.append(ProductLibrary.library_kind == query.library_kind)
    if query.organization_id is not None:
        filters.append(OeOrganization.ocean_account_id == query.organization_id)
        filters.append(OeOrganization.is_deleted == 0)
    joined = (
        select(ProductLibrary, OeOrganization.ocean_account_id)
        .join(OeOrganization, OeOrganization.id == ProductLibrary.organization_id)
        .where(*filters)
    )
    total = int(await session.scalar(select(func.count()).select_from(joined.subquery())) or 0)
    params = PageParams(page=query.page, page_size=query.page_size)
    result = await session.execute(
        joined.order_by(ProductLibrary.id.desc()).offset(params.offset).limit(params.page_size)
    )
    pairs = list(result.all())
    pitchers = await _library_pitchers(session, [row.id for row, _ocean_id in pairs])
    items = [
        _library_data(row, ocean_id, pitchers.get(row.id, []))
        for row, ocean_id in pairs
    ]
    return page_data(items, total, params)


async def create_product_library(session: AsyncSession, body: ProductLibraryBody) -> dict[str, Any]:
    org = await _org_by_ocean_id(session, body.organization_id)
    await _library_no_free(session, body.library_no, None)
    if body.library_role == "fallback":
        await _reject_second_fallback(session, org.id, body.library_kind, None)
    row = ProductLibrary(
        name=body.name,
        library_no=body.library_no,
        library_kind=body.library_kind,
        organization_id=org.id,
        library_role=body.library_role,
        uploaded_count=0,
    )
    session.add(row)
    await session.flush()
    data = _library_data(row, org.ocean_account_id, [])
    await session.commit()
    return data


async def _library_links(session: AsyncSession, library_id: int) -> list[ProductLibraryPitcher]:
    result = await session.execute(
        select(ProductLibraryPitcher).where(ProductLibraryPitcher.product_library_id == library_id)
    )
    return list(result.scalars().all())


async def update_product_library(
    session: AsyncSession, library_id: int, body: ProductLibraryBody
) -> dict[str, Any]:
    row = await _alive_library(session, library_id)
    org = await _org_by_ocean_id(session, body.organization_id)
    links = await _library_links(session, row.id)
    if body.library_role == "fallback" and links:
        raise ApiError(409, "兜底库不能分配投手")
    await _library_no_free(session, body.library_no, row.id)
    if body.library_role == "fallback":
        await _reject_second_fallback(session, org.id, body.library_kind, row.id)
    elif links and (org.id != row.organization_id or body.library_kind != row.library_kind):
        user_ids = [link.user_id for link in links]
        conflict = await session.scalar(
            select(ProductLibraryPitcher.id).where(
                ProductLibraryPitcher.user_id.in_(user_ids),
                ProductLibraryPitcher.organization_id == org.id,
                ProductLibraryPitcher.library_kind == body.library_kind,
                ProductLibraryPitcher.product_library_id != row.id,
            )
        )
        if conflict is not None:
            raise ApiError(409, "投手已挂在同一组织同一类型的其它标准库")
        for link in links:
            link.organization_id = org.id
            link.library_kind = body.library_kind
    row.name = body.name
    row.library_no = body.library_no
    row.library_kind = body.library_kind
    row.organization_id = org.id
    row.library_role = body.library_role
    pitchers = [link.user_id for link in links]
    data = _library_data(row, org.ocean_account_id, pitchers)
    await session.commit()
    return data


async def delete_product_library(session: AsyncSession, library_id: int) -> None:
    row = await _alive_library(session, library_id)
    linked = await _library_pitcher_count(session, row.id)
    if linked or row.uploaded_count > 0:
        raise ApiError(409, "商品库已有投手或已上传短剧")
    row.mark_deleted()
    await session.commit()


async def _library_pitcher_count(session: AsyncSession, library_id: int) -> int:
    total = await session.scalar(
        select(func.count())
        .select_from(ProductLibraryPitcher)
        .where(ProductLibraryPitcher.product_library_id == library_id)
    )
    return int(total or 0)


async def assign_product_library_pitchers(
    session: AsyncSession, library_id: int, body: AssignPitchersBody
) -> dict[str, Any]:
    """覆盖该标准库的投手。别人已挂在同组织同类型的其它库则整批 409。

    从名单去掉的人只删关联，不另插兜底行。兜底库 409。
    """
    row = await _alive_library(session, library_id)
    if row.library_role != "standard":
        raise ApiError(409, "兜底库不能分配投手")
    user_ids = list(dict.fromkeys(body.pitcher_user_ids))
    await _require_pitchers(session, user_ids)
    if user_ids:
        conflict = await session.scalar(
            select(ProductLibraryPitcher.id).where(
                ProductLibraryPitcher.user_id.in_(user_ids),
                ProductLibraryPitcher.organization_id == row.organization_id,
                ProductLibraryPitcher.library_kind == row.library_kind,
                ProductLibraryPitcher.product_library_id != row.id,
            )
        )
        if conflict is not None:
            raise ApiError(409, "投手已挂在同一组织同一类型的其它标准库")
    await session.execute(
        delete(ProductLibraryPitcher).where(ProductLibraryPitcher.product_library_id == row.id)
    )
    await session.flush()
    for user_id in user_ids:
        session.add(
            ProductLibraryPitcher(
                product_library_id=row.id,
                user_id=user_id,
                organization_id=row.organization_id,
                library_kind=row.library_kind,
            )
        )
    await session.commit()
    return {"id": row.id, "pitcher_user_ids": user_ids}


async def sync_advertisers(session: AsyncSession, oe_app_id: int | None) -> dict[str, int]:
    """手动同步广告主。mock 只把种子对齐到表；关闭 mock 时按传入的应用同步。"""
    if get_settings().oceanengine.mock:
        from app.modules.account.seed import ensure_oceanengine_seed

        await ensure_oceanengine_seed(session)
    apps = await _sync_target_apps(session, oe_app_id)
    await _reject_open_advertiser_sync(session, apps)
    runs = await _open_advertiser_runs(session, apps)
    try:
        if not get_settings().oceanengine.mock:
            from app.modules.oceanengine.sync import sync_from_oceanengine

            for app in apps:
                await sync_from_oceanengine(session, app)
        filters = [AdvertiserAccount.is_deleted == 0]
        if oe_app_id is not None:
            filters.append(AdvertiserAccount.oe_app_id == oe_app_id)
        elif apps:
            filters.append(AdvertiserAccount.oe_app_id.in_([app.id for app in apps]))
        total = int(
            await session.scalar(select(func.count()).select_from(AdvertiserAccount).where(*filters)) or 0
        )
        finished = beijing_now()
        for run in runs:
            run.finished_at = finished
            run.success_count = total
            run.failure_count = 0
            run.error_summary = None
        await session.commit()
        return {"advertiser_count": total}
    except Exception as exc:
        await session.rollback()
        finished = beijing_now()
        for run in runs:
            run.finished_at = finished
            run.failure_count = 1
            run.error_summary = str(exc)[:500]
        await session.commit()
        raise


async def _sync_target_apps(session: AsyncSession, oe_app_id: int | None) -> list[OeApp]:
    """指定应用只同步这一套；省略则同步全部有效应用。"""
    if oe_app_id is not None:
        app = await session.get(OeApp, oe_app_id)
        if app is None or app.is_deleted:
            raise ApiError(422, "应用不存在")
        return [app]
    apps = (
        await session.scalars(
            select(OeApp).where(OeApp.is_deleted == 0, OeApp.status == "active").order_by(OeApp.id)
        )
    ).all()
    if not apps:
        raise ApiError(503, "巨量未配置")
    return list(apps)


async def _reject_open_advertiser_sync(session: AsyncSession, apps: list[OeApp]) -> None:
    open_run = await session.scalar(
        select(AccountSyncRun.id).where(
            AccountSyncRun.kind == "advertiser",
            AccountSyncRun.finished_at.is_(None),
            AccountSyncRun.is_deleted == 0,
            AccountSyncRun.oe_app_id.in_([app.id for app in apps]),
        )
    )
    if open_run is not None:
        raise ApiError(409, "已有未结束的广告主同步")


async def _open_advertiser_runs(session: AsyncSession, apps: list[OeApp]) -> list[AccountSyncRun]:
    """先提交进行中的同步行，别的请求才能看见并返回 409。"""
    now = beijing_now()
    runs = [
        AccountSyncRun(oe_app_id=app.id, kind="advertiser", started_at=now, success_count=0, failure_count=0)
        for app in apps
    ]
    session.add_all(runs)
    try:
        await session.commit()
    except IntegrityError as exc:
        await session.rollback()
        raise ApiError(409, "已有未结束的广告主同步") from exc
    return runs


async def choose_library(
    session: AsyncSession,
    pitcher_user_id: int,
    organization_id: int,
    library_kind: str,
) -> int:
    """选上传用的商品库主键。organization_id 是 oe_organization.id，不是巨量账户 id。

    该投手在这个组织、这个类型上有未删除的标准库就用它，否则用兜底库。
    没有兜底库时 409，文案固定为「缺少兜底库」。
    """
    standard_id = await session.scalar(
        select(ProductLibrary.id)
        .join(ProductLibraryPitcher, ProductLibraryPitcher.product_library_id == ProductLibrary.id)
        .where(
            ProductLibraryPitcher.user_id == pitcher_user_id,
            ProductLibraryPitcher.organization_id == organization_id,
            ProductLibraryPitcher.library_kind == library_kind,
            ProductLibrary.is_deleted == 0,
            ProductLibrary.library_role == "standard",
        )
    )
    if standard_id is not None:
        return int(standard_id)
    fallback_id = await session.scalar(
        select(ProductLibrary.id).where(
            ProductLibrary.organization_id == organization_id,
            ProductLibrary.library_kind == library_kind,
            ProductLibrary.library_role == "fallback",
            ProductLibrary.is_deleted == 0,
        )
    )
    if fallback_id is None:
        raise ApiError(409, _MISSING_FALLBACK)
    return int(fallback_id)
