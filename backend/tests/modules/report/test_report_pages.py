"""数据报表：权限、面板拆行、汇总 ROI、趋势小时、固定为 0 的列、组报表数据范围。"""

from __future__ import annotations

import asyncio
import unittest
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any
from unittest.mock import AsyncMock, patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.dialects import postgresql

from app.core.auth import require_token
from app.core.db import get_session
from app.core.envelope import ApiError, register_exception_handlers
from app.core.times import BEIJING
from app.modules.report import crud
from app.modules.report.controller import router
from app.modules.report.export import workbook_bytes
from app.modules.report.pages import uni_speed_data, visible_departments
from app.modules.report.schema import ManhuaDramaQuery, ManhuaProjectQuery, UniSpeedQuery
from app.modules.report.fold import fold_manhua_dramas
from app.modules.report.service import (
    ManhuaFact,
    ProjectLink,
    PromotionLink,
    UniFact,
    arrived_hours,
    manhua_drama_rows,
    manhua_project_rows,
    manhua_promotion_rows,
    manhua_summary,
    manhua_trend,
    uni_drama_rows,
    uni_group_rows,
    uni_pitcher_rows,
    uni_speed,
    uni_summary,
)
from app.modules.system_admin.domain.scope import DataScope

ME = 3
STRANGER = 7
DAY = date(2026, 10, 13)
NOW = datetime(2026, 10, 13, 10, 30, tzinfo=BEIJING)


def link(project_id: int, **kwargs: Any) -> ProjectLink:
    """默认是我（投手 3，部门 10）在剧 8 上建的付费项目。"""
    values: dict[str, Any] = {
        "project_id": project_id,
        "series_id": 8,
        "book_name": "霸总甜宠",
        "pitcher_id": ME,
        "pitcher_name": "甲",
        "department_id": 10,
        "department_name": "投放1组",
        "charge_mode": "IAP",
        "bid_panel": "小额",
        "bid_type": "稳定成本",
        "drama_type": "漫剧",
        "channel": "番茄",
        "publish_date": date(2026, 9, 1),
        "created_at": datetime(2026, 10, 13, 9, 0, tzinfo=BEIJING),
        "project_budget": Decimal("300.00"),
        "roi_goal": Decimal("1.200"),
    }
    values.update(kwargs)
    return ProjectLink(**values)


def fact(project_id: int, promotion_id: int, cost: str, pay: str = "0", **metrics: Any) -> ManhuaFact:
    values = {"stat_cost": Decimal(cost), "stat_pay_amount": Decimal(pay)}
    values.update({key: Decimal(str(value)) for key, value in metrics.items()})
    return ManhuaFact(
        project_id=project_id,
        promotion_id=promotion_id,
        project_name=f"项目{project_id}",
        promotion_name=f"广告{promotion_id}",
        metrics=values,
        first_cost_hour="2026-10-13 08:00:00" if Decimal(cost) > 0 else None,
    )


def iap_query(**kwargs: Any) -> ManhuaDramaQuery:
    return ManhuaDramaQuery(charge_mode="IAP", date_start=DAY, date_end=DAY, **kwargs)


class DramaPermissionTests(unittest.TestCase):
    def test_row_with_me_sums_every_pitcher_and_joins_names(self) -> None:
        links = {
            1: link(1),
            2: link(2, pitcher_id=STRANGER, pitcher_name="乙", department_id=11, department_name="投放2组"),
        }
        facts = [fact(1, 101, "100", "30"), fact(2, 201, "50", "20")]
        rows = manhua_drama_rows(facts, links, iap_query(), ME)
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row["pitchers"], "乙、甲")
        self.assertEqual(row["teams"], "投放1组、投放2组")
        self.assertEqual(row["stat_cost"], 150.0)
        self.assertEqual(row["stat_pay_amount"], 50.0)
        self.assertEqual(row["promotion_with_cost_count"], 2)

    def test_drama_of_only_a_stranger_is_hidden(self) -> None:
        links = {2: link(2, series_id=9, pitcher_id=STRANGER, pitcher_name="乙")}
        rows = manhua_drama_rows([fact(2, 201, "50")], links, iap_query(), ME)
        self.assertEqual(rows, [])

    def test_unmatched_ocean_project_never_shows(self) -> None:
        rows = manhua_drama_rows([fact(999, 9991, "80")], {}, iap_query(), ME)
        self.assertEqual(rows, [])

    def test_project_and_promotion_pages_only_show_my_own(self) -> None:
        links = {1: link(1), 2: link(2, pitcher_id=STRANGER, pitcher_name="乙")}
        facts = [fact(1, 101, "100"), fact(2, 201, "50")]
        query = ManhuaProjectQuery(charge_mode="IAP", date_start=DAY, date_end=DAY)
        projects = manhua_project_rows(facts, links, query, ME)
        self.assertEqual([row["project_id"] for row in projects], [1])
        self.assertEqual(projects[0]["first_cost_time"], "2026-10-13T08:00:00+08:00")
        promotions = manhua_promotion_rows(
            facts, links, {101: PromotionLink(101, 1), 201: PromotionLink(201, 2)}, query, ME
        )
        self.assertEqual([row["promotion_id"] for row in promotions], [101])

    def test_promotion_without_task_row_is_not_listed(self) -> None:
        query = ManhuaProjectQuery(charge_mode="IAP", date_start=DAY, date_end=DAY)
        rows = manhua_promotion_rows([fact(1, 102, "10")], {1: link(1)}, {}, query, ME)
        self.assertEqual(rows, [])


class PanelSplitTests(unittest.TestCase):
    def test_iap_drama_splits_by_panel_and_old_empty_panel_stays_one_row(self) -> None:
        links = {
            1: link(1, bid_panel="小额"),
            2: link(2, bid_panel="大额"),
            3: link(3, bid_panel=""),
        }
        facts = [fact(1, 101, "10"), fact(2, 201, "20"), fact(3, 301, "30"), fact(3, 302, "5")]
        rows = manhua_drama_rows(facts, links, iap_query(), ME)
        by_panel = {row["panel"]: row for row in rows}
        self.assertEqual(set(by_panel), {"小额", "大额", ""})
        self.assertEqual(by_panel[""]["stat_cost"], 35.0)
        self.assertEqual(by_panel["小额"]["stat_cost"], 10.0)

    def test_panel_filter_keeps_only_that_panel(self) -> None:
        links = {1: link(1, bid_panel="小额"), 2: link(2, bid_panel="大额")}
        rows = manhua_drama_rows([fact(1, 101, "10"), fact(2, 201, "20")], links, iap_query(panel=["大额"]), ME)
        self.assertEqual([row["panel"] for row in rows], ["大额"])

    def test_iaa_drama_is_one_row_with_panel_as_column(self) -> None:
        links = {
            1: link(1, charge_mode="IAA", bid_panel=""),
            2: link(2, charge_mode="IAA", bid_panel=""),
        }
        query = ManhuaDramaQuery(charge_mode="IAA", date_start=DAY, date_end=DAY)
        rows = manhua_drama_rows([fact(1, 101, "10"), fact(2, 201, "20")], links, query, ME)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["panel"], "")


class SummaryTests(unittest.TestCase):
    def test_summary_roi_is_sum_of_recovery_over_sum_of_cost(self) -> None:
        links = {1: link(1, series_id=8), 2: link(2, series_id=9)}
        facts = [fact(1, 101, "100", "50"), fact(2, 201, "300", "30")]
        rows = manhua_drama_rows(facts, links, iap_query(), ME)
        self.assertEqual(sorted(row["pay_amount_roi"] for row in rows), [0.1, 0.5])
        summary = manhua_summary(rows)
        self.assertEqual(summary["stat_cost"], 400.0)
        self.assertEqual(summary["pay_amount_roi"], 0.2)

    def test_uni_summary_roi_uses_total_gmv_over_total_cost(self) -> None:
        rows = [
            {"all_gmv_trend": 100, "all_stat_total_cost_trend": 100},
            {"all_gmv_trend": 30, "all_stat_total_cost_trend": 300},
        ]
        self.assertEqual(uni_summary(rows)["all_roi_trend"], 0.325)


class TrendHourTests(unittest.TestCase):
    def test_today_omits_hours_that_have_not_arrived(self) -> None:
        data = manhua_trend(DAY, {3: (Decimal("12.5"), Decimal("4"))}, NOW)
        self.assertEqual([item["hour"] for item in data["items"]], list(range(11)))
        self.assertEqual(data["items"][3], {"hour": 3, "stat_cost": 12.5, "stat_pay_amount": 4.0})
        self.assertEqual(data["items"][0]["stat_cost"], 0.0)

    def test_past_day_is_full_and_future_day_is_empty(self) -> None:
        self.assertEqual(arrived_hours(date(2026, 10, 12), NOW), list(range(24)))
        self.assertEqual(arrived_hours(date(2026, 10, 14), NOW), [])


class UniTests(unittest.TestCase):
    def _facts(self) -> list[UniFact]:
        return [
            UniFact("2026-10-13", 1, metrics={"all_stat_cost_trend": Decimal("100")}),
            UniFact("2026-10-13", 2, metrics={"all_stat_cost_trend": Decimal("40")}),
            UniFact("2026-10-13", 3, metrics={"all_stat_cost_trend": Decimal("7")}),
        ]

    def _links(self) -> dict[int, ProjectLink]:
        return {
            1: link(1, charge_mode=""),
            2: link(2, charge_mode="", pitcher_id=STRANGER, pitcher_name="乙", department_id=11, department_name="投放2组"),
            3: link(3, charge_mode="", pitcher_id=21, pitcher_name="丙", department_id=12, department_name="投放3组"),
        }

    def test_zero_columns_are_zero_and_new_projects_are_counted(self) -> None:
        rows = uni_drama_rows(self._facts(), self._links(), DAY, DAY, ME)
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row["all_stat_cost_trend"], 147.0)
        self.assertEqual(row["new_project_count"], 3)
        for name in ("new_project_with_cost_count", "spending_project_count", "spending_advertiser_count"):
            self.assertEqual(row[name], 0)
        pitchers = uni_pitcher_rows(self._facts(), self._links(), DAY, DAY, ME)
        self.assertEqual(pitchers[0]["new_drama_count"], 0)

    def test_drama_tab_is_one_row_per_drama_per_day_summing_every_pitcher(self) -> None:
        facts = [
            *self._facts()[:2],
            UniFact("2026-10-12", 1, metrics={"all_stat_cost_trend": Decimal("9")}),
        ]
        links = {key: value for key, value in self._links().items() if key != 3}
        rows = uni_drama_rows(facts, links, date(2026, 10, 12), DAY, ME)
        by_day = {row["stat_time_day"]: row for row in rows}
        self.assertEqual(set(by_day), {"2026-10-12", "2026-10-13"})
        self.assertEqual(by_day["2026-10-13"]["all_stat_cost_trend"], 140.0)
        self.assertEqual(by_day["2026-10-12"]["all_stat_cost_trend"], 9.0)
        self.assertNotIn("pitchers", by_day["2026-10-13"])
        self.assertNotIn("pitcher_name", by_day["2026-10-13"])

    def test_pitcher_drama_tab_splits_two_pitchers_on_the_same_day(self) -> None:
        facts = self._facts()[:2]
        links = {key: value for key, value in self._links().items() if key != 3}
        mine = uni_drama_rows(facts, links, DAY, DAY, ME, by_pitcher=True)
        theirs = uni_drama_rows(facts, links, DAY, DAY, STRANGER, by_pitcher=True)
        rows = [*mine, *theirs]
        self.assertEqual(len(rows), 2)
        self.assertEqual(
            sorted((row["pitcher_name"], row["series_id"], row["stat_time_day"], row["all_stat_cost_trend"]) for row in rows),
            [("乙", 8, "2026-10-13", 40.0), ("甲", 8, "2026-10-13", 100.0)],
        )
        self.assertEqual([row["pitcher_user_id"] for row in mine], [ME])
        drama = uni_drama_rows(facts, links, DAY, DAY, ME)
        self.assertEqual(len(drama), 1)
        self.assertEqual(drama[0]["all_stat_cost_trend"], sum(row["all_stat_cost_trend"] for row in rows))

    def test_pitcher_drama_tab_hides_strangers_rows(self) -> None:
        links = {2: self._links()[2]}
        rows = uni_drama_rows(self._facts()[1:2], links, DAY, DAY, ME, by_pitcher=True)
        self.assertEqual(rows, [])

    def test_pitcher_report_shows_only_me_per_day(self) -> None:
        rows = uni_pitcher_rows(self._facts(), self._links(), DAY, DAY, ME)
        self.assertEqual([(row["pitcher_user_id"], row["stat_time_day"]) for row in rows], [(ME, "2026-10-13")])
        self.assertEqual(rows[0]["all_stat_cost_trend"], 100.0)

    def test_group_report_splits_one_drama_into_each_pitchers_department(self) -> None:
        rows = uni_group_rows(self._facts(), self._links(), DAY, DAY, {10, 11})
        by_dept = {row["department_id"]: row["all_stat_cost_trend"] for row in rows}
        self.assertEqual(by_dept, {10: 100.0, 11: 40.0})

    def test_group_scope_falls_back_to_own_department(self) -> None:
        with patch.object(crud, "user_department_id", AsyncMock(return_value=10)):
            empty = asyncio.run(visible_departments(object(), DataScope(ME, True, ())))
            scoped = asyncio.run(visible_departments(object(), DataScope(ME, False, (11, 12))))
        self.assertEqual(empty, {10})
        self.assertEqual(scoped, {11, 12})

    def test_speed_has_growth_rate_and_empty_recovery(self) -> None:
        data = uni_speed(DAY, {0: Decimal("30"), 1: Decimal("5")}, {0: Decimal("20")}, NOW)
        self.assertEqual(data["previous_date"], "2026-10-12")
        self.assertEqual(len(data["cost"]["current"]), 11)
        self.assertEqual(len(data["cost"]["previous"]), 24)
        rates = {item["hour"]: item["rate"] for item in data["cost"]["growth_rate"]}
        self.assertEqual(rates[0], 0.5)
        self.assertEqual(rates[1], 0)
        self.assertEqual(data["recovery"], {"current": [], "previous": [], "growth_rate": []})

    def test_speed_for_a_strangers_pitcher_row_is_not_found(self) -> None:
        query = UniSpeedQuery(scope="pitcher", id=STRANGER, date=DAY)
        with self.assertRaises(ApiError) as caught:
            asyncio.run(uni_speed_data(object(), query, DataScope(ME, True, ()), NOW))
        self.assertEqual(caught.exception.status_code, 404)


class ExportTests(unittest.TestCase):
    def test_workbook_has_chinese_header_and_rows(self) -> None:
        rows = manhua_drama_rows([fact(1, 101, "10", "5")], {1: link(1)}, iap_query(), ME)
        book = load_workbook(BytesIO(workbook_bytes("drama", rows)))
        sheet = book.active
        self.assertEqual(sheet.cell(1, 1).value, "剧目名称")
        self.assertEqual(sheet.cell(2, 1).value, "霸总甜宠")
        self.assertEqual(sheet.max_row, 2)


class _Result:
    def all(self) -> list[Any]:
        return []


class _SqlSession:
    """记下语句，按 PostgreSQL 编译一遍，返回空结果。"""

    def __init__(self) -> None:
        self.sql: list[str] = []

    async def execute(self, statement: Any) -> _Result:
        self.sql.append(str(statement.compile(dialect=postgresql.dialect())))
        return _Result()

    async def scalar(self, statement: Any) -> Any:
        self.sql.append(str(statement.compile(dialect=postgresql.dialect())))
        return None


class ReadOnlySnapshotTests(unittest.TestCase):
    def test_loaders_only_read_snapshot_and_task_tables(self) -> None:
        session = _SqlSession()

        async def run() -> None:
            await crud.manhua_facts(session, DAY, DAY, "IAP")
            await crud.standard_links(session, charge_mode="IAP", project_ids=[1])
            await crud.standard_promotions(session, [101])
            await crud.manhua_first_cost(session, [1], by_promotion=False)
            await crud.manhua_hourly(session, DAY, project_ids=[1])
            await crud.uni_facts(session, DAY, DAY)
            await crud.uni_links(session, project_ids=[1], created_between=(DAY, DAY))
            await crud.uni_hourly_cost(session, DAY, [1])

        asyncio.run(run())
        text = "\n".join(session.sql)
        self.assertIn("oe_delivery_report", text)
        self.assertIn("oe_uni_project_report", text)
        self.assertIn("delivery_task_project", text)
        self.assertIn("uni_native_task_account", text)
        self.assertTrue(all(sql.lstrip().upper().startswith("SELECT") for sql in session.sql))


def _client(menus: list[str]) -> TestClient:
    app = FastAPI()
    register_exception_handlers(app)
    app.include_router(router)

    async def _session():
        yield object()

    app.dependency_overrides[get_session] = _session
    app.dependency_overrides[require_token] = lambda: {
        "id": str(ME),
        "nickname": "甲",
        "login_account": "pitcher",
        "tenant": "agent",
        "enabled": True,
        "menu_ids": menus,
        "data_scope": {"self_only": True, "department_ids": []},
    }
    return TestClient(app, raise_server_exceptions=False)


class HttpTests(unittest.TestCase):
    def test_iap_list_needs_menu_9(self) -> None:
        response = _client(["5"]).get("/api/v1/reports/manhua/dramas", params={"charge_mode": "IAP"})
        self.assertEqual(response.status_code, 403)

    def test_iaa_rejects_panel_filter(self) -> None:
        response = _client(["5"]).get(
            "/api/v1/reports/manhua/projects", params={"charge_mode": "IAA", "panel": "小额"}
        )
        self.assertEqual(response.status_code, 422)

    def test_iap_drama_rejects_filters_outside_its_prototype(self) -> None:
        response = _client(["9"]).get(
            "/api/v1/reports/manhua/dramas", params={"charge_mode": "IAP", "series_name": "霸总"}
        )
        self.assertEqual(response.status_code, 422)

    def test_drama_list_pages_with_summary_and_never_calls_ocean(self) -> None:
        facts = [fact(1, 101, "100", "50"), fact(2, 201, "300", "30")]
        links = {1: link(1, series_id=8), 2: link(2, series_id=9)}
        built = manhua_drama_rows(facts, links, iap_query(), ME)
        with (
            patch("app.modules.report.pages.read_serving.manhua_drama_rows", AsyncMock(return_value=built)),
            patch("app.modules.oceanengine.client.OceanEngineClient.custom_report") as ocean,
        ):
            response = _client(["9"]).get(
                "/api/v1/reports/manhua/dramas",
                params={"charge_mode": "IAP", "page_size": 1, "date_start": "2026-10-13", "date_end": "2026-10-13"},
            )
        ocean.assert_not_called()
        self.assertEqual(response.status_code, 200, response.text)
        data = response.json()["data"]
        self.assertEqual(data["total"], 2)
        self.assertEqual(len(data["list"]), 1)
        self.assertEqual(data["list"][0]["stat_cost"], 300.0)
        self.assertEqual(data["summary"]["pay_amount_roi"], 0.2)

    def test_group_report_uses_data_scope(self) -> None:
        facts = [
            UniFact("2026-10-13", 1, metrics={"all_stat_cost_trend": Decimal("100")}),
            UniFact("2026-10-13", 2, metrics={"all_stat_cost_trend": Decimal("40")}),
        ]
        links = {
            1: link(1),
            2: link(2, pitcher_id=STRANGER, pitcher_name="乙", department_id=11, department_name="投放2组"),
        }
        built = uni_group_rows(facts, links, DAY, DAY, {11})
        reader = AsyncMock(return_value=built)
        with (
            patch("app.modules.report.pages.read_serving.uni_group_rows", reader),
            patch.object(crud, "user_department_id", AsyncMock(return_value=11)),
        ):
            response = _client(["15"]).get(
                "/api/v1/reports/uni/groups", params={"date_start": "2026-10-13", "date_end": "2026-10-13"}
            )
        self.assertEqual(reader.await_args.args[2], {11})
        self.assertEqual(response.status_code, 200, response.text)
        rows = response.json()["data"]["list"]
        self.assertEqual([(row["department_id"], row["all_stat_cost_trend"]) for row in rows], [(11, 40.0)])


class FoldTests(unittest.TestCase):
    def test_two_days_of_one_drama_sum_and_count_an_ad_once(self) -> None:
        base = {
            "series_id": 8,
            "panel_key": "小额",
            "book_name": "霸总甜宠",
            "drama_type": "漫剧",
            "channel": "番茄",
            "publish_date": date(2026, 9, 1),
            "pitcher_ids": [ME],
            "pitcher_names": ["甲"],
            "department_names": ["投放1组"],
            "bid_types": ["稳定成本"],
            "panel_names": ["小额"],
            "costed_promotion_ids": [101],
        }
        rows = [
            {**base, "stat_cost": Decimal("100"), "stat_pay_amount": Decimal("30")},
            {**base, "stat_cost": Decimal("50"), "stat_pay_amount": Decimal("20")},
        ]
        query = ManhuaDramaQuery(
            charge_mode="IAP", date_start=date(2026, 10, 12), date_end=date(2026, 10, 13)
        )
        folded = fold_manhua_dramas(rows, query, ME)
        self.assertEqual(len(folded), 1)
        self.assertEqual(folded[0]["stat_cost"], 150.0)
        self.assertEqual(folded[0]["promotion_with_cost_count"], 1)
        self.assertEqual(folded[0]["pay_amount_roi"], 0.3333)

    def test_owner_keeps_a_row_that_is_not_theirs(self) -> None:
        row = {
            "series_id": 8,
            "panel_key": "小额",
            "book_name": "霸总甜宠",
            "drama_type": "漫剧",
            "channel": "番茄",
            "publish_date": date(2026, 9, 1),
            "pitcher_ids": [STRANGER],
            "pitcher_names": ["乙"],
            "department_names": ["投放2组"],
            "bid_types": ["稳定成本"],
            "panel_names": ["小额"],
            "costed_promotion_ids": [],
            "stat_cost": Decimal("40"),
            "stat_pay_amount": Decimal("10"),
        }
        folded = fold_manhua_dramas([row], iap_query(), None)
        self.assertEqual(len(folded), 1)
        self.assertEqual(folded[0]["pitchers"], "乙")


if __name__ == "__main__":
    unittest.main()
