"""漫剧列表的时间展示和筛选条件。"""

from __future__ import annotations

import io
import unittest
from datetime import datetime
from typing import Annotated, get_args
from unittest.mock import AsyncMock, patch

from fastapi import FastAPI, Query
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from app.core.db import get_session
from app.core.times import BEIJING
from app.modules.material.controller import PrincipalDep, router
from app.modules.material.schema import ManhuaSeriesQuery
from app.modules.material.service import export_cells, rows_to_xlsx, to_item


class StampParamTests(unittest.TestCase):
    def test_query_rejects_date_without_time(self) -> None:
        """入参不是 YYYY-MM-DD HH:MM:SS 时，请求在进业务前被拒绝。"""
        app = FastAPI()

        @app.get("/t")
        def _probe(query: Annotated[ManhuaSeriesQuery, Query()]) -> dict[str, datetime | None]:
            """试查询模型的时间入参。"""
            return {"value": query.estimate_publish_time_from}

        client = TestClient(app)
        rejected = client.get("/t", params={"estimate_publish_time_from": "2026-09-22"})
        self.assertEqual(rejected.status_code, 422)
        accepted = client.get("/t", params={"estimate_publish_time_from": "2026-09-22 10:00:00"})
        self.assertEqual(accepted.status_code, 200)
        self.assertEqual(accepted.json()["value"], "2026-09-22T10:00:00+08:00")


class QueryRangeTests(unittest.TestCase):
    def test_reversed_time_and_episodes_are_rejected(self) -> None:
        """结束早于开始、集数上限小于下限，入参直接拒绝。"""
        with self.assertRaises(ValidationError):
            ManhuaSeriesQuery(
                estimate_publish_time_from="2026-09-22 10:00:00",
                estimate_publish_time_to="2026-09-21 10:00:00",
            )
        with self.assertRaises(ValidationError):
            ManhuaSeriesQuery(episode_amount_min=10, episode_amount_max=2)


class ItemTests(unittest.TestCase):
    def test_item_returns_tab_text(self) -> None:
        """列表项原样返回 IAA/IAP，部门恒空。"""

        class Row:
            id = 7
            playlet_id = 9
            book_id = 8
            category_text = "玄幻脑洞"
            tab_text = "IAP"
            single_price = "50"
            thumb_url = ""
            book_name = "剧"
            episode_amount = 12
            publish_status = 2
            delivery_status = True
            publish_time = "2026-09-22 01:00:00"
            estimate_publish_time = datetime(2026, 9, 21, 17, 30, tzinfo=BEIJING)
            create_time = "坏的"
            collected_at = datetime(2026, 9, 22, 1, 0, tzinfo=BEIJING)
            douyin_nick_name = "号"

        item = to_item(Row())
        self.assertEqual(item["tab_text"], "IAP")
        self.assertEqual(item["publish_time"], "2026-09-22 01:00:00")
        self.assertEqual(item["estimate_publish_time"], "2026-09-21 17:30:00")
        self.assertEqual(item["create_time"], "坏的")
        self.assertIsNone(item["department_name"])
        self.assertEqual(item["id"], "7")
        self.assertEqual(item["collected_at"], "2026-09-22T01:00:00+08:00")


class ExportTests(unittest.TestCase):
    def test_department_query_is_rejected(self) -> None:
        """查询不再收部门。"""
        with self.assertRaises(ValidationError):
            ManhuaSeriesQuery(department_id="9")

    def test_xlsx_keeps_ids_and_status_text(self) -> None:
        """导出把发布状态写成中文，专辑 ID 按文本留下。"""

        class Row:
            id = 7
            playlet_id = 1234567890123456789
            book_id = 8
            category_text = "玄幻"
            tab_text = "IAA"
            thumb_url = ""
            book_name = "剧名"
            episode_amount = 12
            publish_status = 2
            delivery_status = False
            publish_time = ""
            estimate_publish_time = None
            create_time = ""
            collected_at = datetime(2026, 9, 22, 1, 0, tzinfo=BEIJING)
            douyin_nick_name = "号"

        cells = export_cells(to_item(Row()))
        self.assertEqual(cells[0], "剧名")
        self.assertEqual(cells[2], "已发布")
        self.assertEqual(cells[4], "否")
        self.assertEqual(cells[7], "2026-09-22 01:00:00")
        self.assertEqual(cells[8], "1234567890123456789")
        self.assertEqual(cells[9], "号")
        self.assertNotIn("部门", cells)
        today = datetime.now(BEIJING).strftime("%Y-%m-%d 12:00:00")
        listed = export_cells({"publish_time": today, "publish_status": 1, "collected_at": ""})
        self.assertEqual(listed[4], "是")
        payload = rows_to_xlsx(["短剧名称", "抖音id"], [["剧名", cells[8]]])
        sheet = load_workbook(io.BytesIO(payload)).active
        assert sheet is not None
        self.assertEqual(sheet["B2"].value, "1234567890123456789")

    def test_export_route_returns_xlsx(self) -> None:
        """导出走查询参数，成功时直接下 xlsx，不包信封。"""
        app = FastAPI()
        app.include_router(router)

        async def _session():
            yield None

        principal = get_args(PrincipalDep)[1].dependency
        app.dependency_overrides[get_session] = _session
        app.dependency_overrides[principal] = lambda: {"menu_ids": ["95"], "enabled": True}
        client = TestClient(app)
        with patch("app.modules.material.service.list_series", new=AsyncMock(return_value=[])):
            res = client.post("/api/v1/material/manhua-series/export")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            res.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("manhua_series.xlsx", res.headers["content-disposition"])
        sheet = load_workbook(io.BytesIO(res.content)).active
        assert sheet is not None
        self.assertEqual(sheet["A1"].value, "短剧名称")
        self.assertEqual(sheet["E1"].value, "是否当天上架")
        self.assertEqual(sheet["I1"].value, "抖音id")
        self.assertEqual(sheet["J1"].value, "抖音名")
        self.assertEqual(sheet.max_column, 10)
        self.assertEqual(sheet.max_row, 1)
